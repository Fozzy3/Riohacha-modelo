"""Parse 'Pto CIP Guajira 5.xlsx' into structured JSON for the BIM model.

Usage: uv run --with openpyxl python -I tools/parse_presupuesto.py <xlsx> <memoria_indice.json> <out_dir>
Writes:
  presupuesto.json  infrastructure items (sheet 08) with chapter, AIU block and per-floor areas
  dotacion.json     equipment per lab (sheet 06)
  mobiliario.json   furniture counts per space and floor (sheet 03.1)
and prints a cross-check of quantities against the memoria de cantidades.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ITEM = re.compile(r"^\d{1,2}\.\d{2,3}(\.\d{2})?$")  # 1.01, 9.100, 12.01.02
CHAPTER = re.compile(r"^\s*(\d{1,2})\.00(?:\.01)?\s+(.+)$")


def s(v):
    return str(v).strip() if v is not None else ""


def infraestructura(ws):
    items, chapters, aiu, areas = [], {}, {}, {}
    chapter, group = None, ""
    for r in ws.iter_rows(values_only=True):
        b, c = s(r[1]), r[2]
        if re.match(r"^\d{1,2}\.\d$", b):
            b += "0"  # Excel stored 10.10 as the number 10.1
        m = CHAPTER.match(b)
        if m and not ITEM.match(b):
            if chapter != int(m.group(1)):
                group = ""
            chapter = int(m.group(1))
            chapters.setdefault(chapter, m.group(2).strip())
            continue
        if ITEM.match(b) and s(r[3]) == "U.M.":
            continue  # sub-header row of a 3-level chapter (12.01.01 ... U.M. Cantidad ...)
        if ITEM.match(b) and not isinstance(r[4], (int, float)):
            group = s(c)  # subgroup title inside a chapter, e.g. 9.26 RED DE VENTILACIÓN
            continue
        if ITEM.match(b):
            # Chapter comes from the section the row sits in; some codes are mistyped (29.10 inside 28).
            items.append({"item": b, "capitulo": chapter, "subgrupo": group, "descripcion": s(c),
                          "unidad": s(r[3]), "cantidad": r[4], "vr_unitario": r[5], "vr_total": r[6]})
            continue
        if b.upper() in {"ADMINISTRACION", "ADMINISTRACIÓN", "IMPREVISTOS", "UTILIDAD", "IVA SOBRE UTILIDAD"}:
            aiu[b.upper()] = {"porcentaje": r[2], "valor": r[3]}
        for k in range(1, 6):
            if re.match(r"^\d(ER|DO|TO) PISO", s(r[k]).upper()):
                areas[s(r[k]).upper()] = s(r[k + 1])
    for it in items:
        it["capitulo_nombre"] = chapters.get(it["capitulo"], "")
    return items, chapters, aiu, areas


def dotacion(ws):
    labs, cur = [], None
    for r in ws.iter_rows(values_only=True):
        b = s(r[1])
        if b and all(x in (None, "") for x in r[2:10]) and not b.startswith(("Total", "Rubros", "Valores")):
            cur = {"espacio": b, "equipos": []}
            labs.append(cur)
            continue
        if cur and r[2] and isinstance(r[5], (int, float)) and isinstance(r[8], (int, float)):
            cur["equipos"].append({"equipo": s(r[2]), "especificaciones": s(r[3]), "justificacion": s(r[4]),
                                   "cantidad": r[5], "unidad": s(r[6]), "custodia": s(r[7]),
                                   "vr_unitario": r[8], "vr_total": r[9]})
    return [lab for lab in labs if lab["equipos"]]


def mobiliario(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if s(r[1]) == "Espacio")
    kinds = {k: s(v) for k, v in enumerate(rows[header]) if k >= 5 and v}
    spaces, piso = [], None
    for r in rows[header + 1:]:
        if s(r[1]).lower().startswith("piso"):
            piso = s(r[1]).capitalize()  # the sheet writes 'piso 2' once in lower case
        name = s(r[3])
        if piso and name and isinstance(r[2], (int, float)):
            counts = {kinds[k]: r[k] for k in kinds if isinstance(r[k], (int, float)) and r[k]}
            spaces.append({"piso": piso, "numero": int(r[2]), "espacio": name,
                           "cantidad_espacios": r[4] if isinstance(r[4], (int, float)) else None,
                           "mobiliario": counts})
        if s(r[1]) == "Total":
            break
    return spaces


def main(xlsx, memoria, out_dir):
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    items, chapters, aiu, areas = infraestructura(wb["08. Infraestructura"])
    (out / "presupuesto.json").write_text(json.dumps(
        {"capitulos": chapters, "items": items, "aiu": aiu, "areas_por_piso": areas}, ensure_ascii=False, indent=1))
    labs = dotacion(wb["06. Equipos y Software"])
    (out / "dotacion.json").write_text(json.dumps(labs, ensure_ascii=False, indent=1))
    spaces = mobiliario(wb["03.1 Mobiliario"])
    (out / "mobiliario.json").write_text(json.dumps(spaces, ensure_ascii=False, indent=1))

    total = sum(i["vr_total"] or 0 for i in items)
    print(f"presupuesto: {len(items)} items, {len(chapters)} capitulos, suma items {total:,.0f}"
          f" (AIU indica costo directo 13,288,996,949 sin seguridad)")
    by_ch = {}
    for i in items:
        by_ch[i["capitulo"]] = by_ch.get(i["capitulo"], 0) + (i["vr_total"] or 0)
    print("  por capitulo:", {k: round(v / 1e6, 1) for k, v in sorted(by_ch.items())})
    print(f"dotacion: {len(labs)} grupos, {sum(len(l['equipos']) for l in labs)} equipos")
    print(f"mobiliario: {len(spaces)} espacios")
    print("areas:", areas)

    mem = {i["item"]: i for i in json.load(open(memoria))}
    missing = [i["item"] for i in items if i["item"] not in mem]
    diff = [(i["item"], i["cantidad"], mem[i["item"]]["cantidad"]) for i in items
            if i["item"] in mem and abs(i["cantidad"] - (mem[i["item"]]["cantidad"] or 0)) > max(0.01, 0.005 * abs(i["cantidad"]))]
    print(f"cruce memoria: {len(items) - len(missing)} con memoria, {len(missing)} sin memoria, {len(diff)} cantidades distintas")
    print("  sin memoria:", missing[:60])
    print("  distintas (item, presupuesto, memoria):", diff[:30])


if __name__ == "__main__":
    main(*sys.argv[1:])
